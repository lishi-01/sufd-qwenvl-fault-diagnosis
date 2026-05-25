from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT = PROJECT_ROOT / "docs" / "SUFD_QwenVL_RLHF_project_plan.docx"

TITLE = "SUFD-QwenVL 工业故障诊断强化学习阶段执行方案"
SUBTITLE = "扩展 DPO、PPO、GRPO 与奖励模型训练流程"

BLUE = "1F4E79"
ACCENT = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "F4F6F8"
BORDER = "A6A6A6"


def qfont(run, size=None, bold=None, color=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        node = tc_borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, size=9.2, fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    borders(cell)
    margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    qfont(run, size=size, bold=bold)


def para(doc, text="", style=None, size=None, bold=False, color=None, align=None, before=0, after=5):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        qfont(r, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        qfont(r, color=BLUE if level <= 2 else ACCENT)
    return p


def bullets(doc, items):
    for item in items:
        para(doc, item, style="List Bullet", after=2)


def table(doc, headers, rows, widths=None, header_fill=LIGHT_BLUE, size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for idx, h in enumerate(headers):
        cell_text(t.rows[0].cells[idx], h, bold=True, size=size, fill=header_fill, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(val)) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[idx], val, size=size, align=align)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    para(doc, after=3)
    return t


def code(doc, text, title=None):
    if title:
        para(doc, title, bold=True, color=BLUE, after=2)
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, "F7F7F7")
    borders(c, color="D0D0D0")
    margins(c, top=110, start=130, bottom=110, end=130)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(text.strip("\n").split("\n")):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        qfont(r, size=8.4, font="Consolas")
    para(doc, after=3)


def callout(doc, title, body, fill=LIGHT_BLUE):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, fill)
    borders(c, color="8EAADB")
    margins(c, top=140, start=160, bottom=140, end=160)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    qfont(r, size=10.2, bold=True, color=BLUE)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    qfont(r2, size=9.2)
    para(doc, after=2)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10)
    for style_name, size in [("Heading 1", 15.5), ("Heading 2", 12.5), ("Heading 3", 11)]:
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True
    return doc


def build():
    doc = setup_doc()

    para(doc, TITLE, size=23, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=70, after=8)
    para(doc, SUBTITLE, size=14.5, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    para(doc, "基于当前 Qwen2.5-VL-3B + LoRA SFT 项目计划与实验结果扩展", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

    table(doc, ["项目项", "当前设定"], [
        ["项目基线", "Qwen2.5-VL-3B-Instruct + LoRA，class_only，checkpoint-2200"],
        ["当前最佳指标", "Accuracy 0.4854，Macro-F1 0.4805，格式正确率 1.0000"],
        ["RL 阶段目标", "提升回答偏好、格式稳定性、诊断一致性与报告可信度"],
        ["推荐优先级", "DPO → 奖励模型 RM → PPO → GRPO 对照实验"],
        ["文档版本", "v0.1，2026-05-24"],
    ], widths=[3.2, 12.5], size=9.4)

    callout(doc, "核心判断", "强化学习不建议用来直接“硬救”当前五分类准确率。它更适合在已有 SFT 分类能力基础上，优化回答偏好、降低幻觉、强化格式约束，并提升诊断依据与维护建议的一致性。分类性能仍应主要通过图像表示、模型容量、验证集与 SFT 参数继续提升。", LIGHT_ORANGE)

    heading(doc, "目录", 1)
    for item in [
        "1. RL 阶段定位与总体路线",
        "2. 数据与奖励设计总览",
        "3. 偏好数据构造规范",
        "4. DPO 训练流程",
        "5. 奖励模型训练流程",
        "6. PPO 训练流程",
        "7. GRPO 训练流程",
        "8. 评估指标与实验矩阵",
        "9. 工程脚本与交付物规划",
        "10. 风险控制与推荐执行顺序",
    ]:
        para(doc, item, after=1)

    doc.add_page_break()

    heading(doc, "1. RL 阶段定位与总体路线", 1)
    para(doc, "当前项目已经完成从 SUFD 原始振动信号到包络 STFT 图像、多模态 SFT 数据、Qwen2.5-VL LoRA 微调、测试评估和 Gradio Demo 的完整闭环。RL 阶段应作为“对齐与报告质量增强层”，而不是替代现有 SFT 分类模型。")
    table(doc, ["阶段", "输入资产", "训练目标", "主要产出"], [
        ["SFT 基线", "train_sft_class_only.json、包络 STFT 图像、工况文本", "学习五分类故障类型输出", "checkpoint-2200 最佳 class_only 模型"],
        ["DPO", "同一 prompt 下 chosen/rejected 回答对", "偏好正确标签、稳定格式和一致诊断文本", "DPO adapter、偏好胜率评估"],
        ["奖励模型 RM", "chosen/rejected 偏好对", "学习回答质量打分函数", "Reward Model、RM pairwise accuracy"],
        ["PPO", "SFT/DPO policy + RM + 参考模型", "在线生成并最大化奖励，同时用 KL 约束防止漂移", "PPO adapter、报告质量提升"],
        ["GRPO", "每个 prompt 生成多个候选回答 + 规则/RM 奖励", "用组内相对优势优化，无需单独 value model", "GRPO adapter、与 DPO/PPO 对照"],
    ], widths=[2.2, 4.5, 5.0, 4.0], size=8.4)
    code(doc, """现有最佳 SFT class_only 模型 checkpoint-2200
        ↓
构造偏好数据：chosen / rejected
        ↓
DPO：低成本偏好优化，优先跑通
        ↓
训练奖励模型 RM：学习诊断回答质量评分
        ↓
PPO：基于 RM 在线优化完整回答
        ↓
GRPO：基于多候选组内相对奖励进行对照实验
        ↓
统一评估：分类指标 + 格式正确率 + 一致性 + 幻觉率 + 人工抽检""", "推荐总流程")
    bullets(doc, [
        "让模型更偏向“标签正确、格式正确、依据与标签一致、维护建议合理”的回答。",
        "降低幻觉：不生成不存在的故障类型、不编造未提供的传感器或工况。",
        "让诊断依据和维护建议更稳定，不因提示词轻微变化而漂移。",
        "为后续 Demo 和简历项目增加“偏好对齐 / RLHF”能力展示。",
    ])

    heading(doc, "2. 数据与奖励设计总览", 1)
    para(doc, "RL 阶段的核心不是重新生成图像，而是围绕同一张图像、同一工况和同一 prompt 生成多个回答，并建立“好回答”和“差回答”的偏好关系。")
    table(doc, ["质量维度", "正向标准", "负向样例", "可自动评分"], [
        ["故障类型正确性", "输出标签与 metadata 的 label 一致", "真实 BF 却输出正常状态", "是"],
        ["格式稳定性", "字段完整，可稳定抽取故障类型", "漏字段、字段顺序混乱", "是"],
        ["诊断一致性", "诊断依据与故障类型一致", "类型为内圈，依据描述外圈剥落", "部分自动"],
        ["维护建议一致性", "建议检查部位与预测故障匹配", "滚动体故障却建议只检查电机转矩", "部分自动"],
        ["工况忠实性", "只使用输入中的转速、负载、采样率", "编造 50 Hz、5 V 或其他传感器", "是"],
        ["表达质量", "简洁、专业、无无关内容", "长篇泛化、重复、模糊措辞", "需 RM/人工"],
    ], widths=[3.0, 4.9, 5.0, 2.0], size=8.3)
    code(doc, """reward = 0.45 * label_reward
       + 0.20 * format_reward
       + 0.15 * consistency_reward
       + 0.10 * condition_faithfulness_reward
       + 0.10 * brevity_and_style_reward

label_reward: 标签正确为 1，否则 0
format_reward: 字段完整且可解析为 1，否则 0
consistency_reward: 依据/建议与故障类型匹配为 1，否则扣分
condition_faithfulness_reward: 不编造工况为 1，否则扣分
brevity_and_style_reward: 输出简洁、专业、无重复为正分""", "多目标奖励建议")
    callout(doc, "建议", "DPO 阶段优先使用“硬偏好”：正确标签 + 正确格式作为 chosen；错误标签或错误格式作为 rejected。PPO/GRPO 阶段再引入连续奖励，避免一开始奖励设计过复杂。", LIGHT_GREEN)

    heading(doc, "3. 偏好数据构造规范", 1)
    para(doc, "LLaMA-Factory 的偏好数据可以使用 ShareGPT 格式，并提供 chosen 和 rejected 字段。该数据可同时用于 DPO、奖励模型训练、ORPO/SimPO 等偏好优化任务。")
    table(doc, ["来源", "chosen 构造方式", "rejected 构造方式", "适用阶段"], [
        ["规则模板", "真实 label 对应的正确完整报告", "错误 label 的报告、字段缺失报告、错工况报告", "DPO / RM"],
        ["模型生成", "checkpoint-2200 预测正确且格式正确的回答", "预测错误、格式异常或依据不一致的回答", "DPO / RM"],
        ["人工审核", "人工认为更专业的一版报告", "人工认为不可信或不一致的一版报告", "RM / DPO"],
        ["多候选采样", "同一 prompt 采样 N 个回答后按规则/RM 选最高分", "同组最低分回答", "DPO / GRPO"],
    ], widths=[2.7, 4.7, 5.1, 2.4], size=8.3)
    code(doc, """{
  \"conversations\": [
    {\"from\": \"human\", \"value\": \"<image>\\n你是一名工业设备故障诊断助手...\"}
  ],
  \"chosen\": {
    \"from\": \"gpt\",
    \"value\": \"故障类型：滚动体故障\\n诊断依据：...\\n维护建议：...\"
  },
  \"rejected\": {
    \"from\": \"gpt\",
    \"value\": \"故障类型：正常状态\\n诊断依据：...\\n维护建议：...\"
  },
  \"images\": [\"data/processed/images/all/sample_BF_20_0_000409.png\"]
}""", "ShareGPT 偏好样本示例")
    code(doc, """\"sufd_fault_dpo_train\": {
  \"file_name\": \"sufd/train_dpo.json\",
  \"formatting\": \"sharegpt\",
  \"ranking\": true,
  \"columns\": {
    \"messages\": \"conversations\",
    \"chosen\": \"chosen\",
    \"rejected\": \"rejected\",
    \"images\": \"images\"
  }
}""", "dataset_info.json 注册示例")

    heading(doc, "4. DPO 训练流程", 1)
    para(doc, "DPO 是本项目最推荐优先实现的偏好优化阶段。它不需要单独训练奖励模型，也不需要在线采样更新，工程复杂度明显低于 PPO。DPO 的目标是让模型更偏好 chosen 回答，而远离 rejected 回答。")
    table(doc, ["项目", "建议设置"], [
        ["初始模型", "checkpoint-2200 class_only adapter；也可对比 stage2 full adapter"],
        ["数据", "train_dpo.json，包含 chosen/rejected/images"],
        ["训练目标", "保持故障类型正确，同时学习完整诊断报告偏好"],
        ["关键超参", "pref_beta 0.05~0.2；learning_rate 1e-5~5e-5；epoch 1~3"],
        ["主要风险", "偏好数据质量差会拉偏模型；过大 beta/学习率会损伤分类能力"],
    ], widths=[3.0, 12.0], size=8.8)
    code(doc, """### model
model_name_or_path: /root/autodl-tmp/models/Qwen2.5-VL-3B-Instruct
adapter_name_or_path: saves/qwen2_5vl_3b/lora/sufd_class_only_r16_ep5/checkpoint-2200
trust_remote_code: true

### method
stage: dpo
do_train: true
finetuning_type: lora
lora_target: all
lora_rank: 16
lora_alpha: 32
lora_dropout: 0.05
pref_beta: 0.1
pref_loss: sigmoid

### dataset
dataset: sufd_fault_dpo_train
template: qwen2_vl
cutoff_len: 2048
overwrite_cache: true
preprocessing_num_workers: 4

### train
per_device_train_batch_size: 2
gradient_accumulation_steps: 4
learning_rate: 3.0e-5
num_train_epochs: 2.0
bf16: true""", "DPO 配置模板")
    bullets(doc, [
        "分类指标：Accuracy、Macro-F1 不能明显低于 checkpoint-2200。",
        "格式正确率：应保持 1.0000 或接近 1.0000。",
        "偏好胜率：同一 prompt 下，DPO 回答优于 SFT 回答的比例。",
        "一致性指标：故障类型与诊断依据、维护建议是否匹配。",
    ])

    heading(doc, "5. 奖励模型训练流程", 1)
    para(doc, "奖励模型用于给回答质量打分，是 PPO 的核心组件，也可以作为 GRPO 的 reward source 之一。奖励模型训练仍使用 chosen/rejected 偏好对，但输出不是文本，而是一个标量 reward。")
    table(doc, ["模块", "说明"], [
        ["输入", "prompt + image + assistant response"],
        ["监督信号", "chosen 的 reward 应高于 rejected"],
        ["训练目标", "pairwise ranking loss，使 RM 能区分高质量/低质量诊断回答"],
        ["验证指标", "RM pairwise accuracy，即 RM 给 chosen 更高分的比例"],
        ["上线用途", "PPO reward、GRPO reward、best-of-N 选择、报告质量审核"],
    ], widths=[3.0, 12.0], size=8.8)
    code(doc, """### method
stage: rm
do_train: true
finetuning_type: lora
lora_target: all
lora_rank: 16
lora_alpha: 32

### dataset
dataset: sufd_fault_rm_train
template: qwen2_vl
cutoff_len: 2048

### output
output_dir: saves/qwen2_5vl_3b/reward/sufd_rm

### train
per_device_train_batch_size: 2
gradient_accumulation_steps: 4
learning_rate: 2.0e-5
num_train_epochs: 2.0
bf16: true""", "奖励模型配置模板")
    bullets(doc, [
        "chosen/rejected 差异必须明确，避免两个回答都很好或都很差。",
        "优先构造“标签正确 vs 标签错误”“格式完整 vs 格式缺失”“依据一致 vs 依据错配”的强偏好。",
        "RM 验证集必须独立，不能只看训练集 pairwise accuracy。",
        "如果 RM 对错误标签仍给高分，则不能进入 PPO。",
    ])

    heading(doc, "6. PPO 训练流程", 1)
    para(doc, "PPO 是经典 RLHF 流程：策略模型生成回答，奖励模型打分，再通过 PPO 更新策略，同时用 KL 惩罚防止模型偏离参考模型。它比 DPO 更灵活，但工程成本和显存成本更高。")
    code(doc, """SFT/DPO policy 生成回答
        ↓
Reward Model 评分
        ↓
KL(policy || reference) 约束
        ↓
PPO 更新 LoRA policy
        ↓
评估分类、格式、诊断一致性和奖励分布""", "PPO 闭环")
    table(doc, ["组件", "本项目建议"], [
        ["Policy 初始值", "checkpoint-2200 或 DPO adapter"],
        ["Reference model", "冻结的 checkpoint-2200 或同基座模型"],
        ["Reward model", "第 5 节训练的 sufd_rm"],
        ["Reward 组合", "RM 分数 + 规则奖励，如标签正确、格式正确、工况忠实"],
        ["KL 控制", "必须开启，避免模型为追求奖励而输出怪异文本"],
    ], widths=[3.0, 12.0], size=8.8)
    code(doc, """### method
stage: ppo
do_train: true
finetuning_type: lora

### model
model_name_or_path: /root/autodl-tmp/models/Qwen2.5-VL-3B-Instruct
adapter_name_or_path: saves/qwen2_5vl_3b/lora/sufd_dpo_from_best
reward_model: saves/qwen2_5vl_3b/reward/sufd_rm
reward_model_type: lora

### train
per_device_train_batch_size: 1
gradient_accumulation_steps: 8
learning_rate: 1.0e-6
num_train_epochs: 1.0
bf16: true

### generation
generation_max_length: 256
top_p: 0.9
temperature: 0.7""", "PPO 配置模板")
    callout(doc, "PPO 注意事项", "PPO 学习率通常要比 SFT/DPO 小一个数量级。建议先小样本跑 50~100 条 prompt，观察 reward、KL、长度和格式，再进行全量训练。若分类准确率明显下降，应立即降低学习率或增大 KL 约束。", LIGHT_ORANGE)

    heading(doc, "7. GRPO 训练流程", 1)
    para(doc, "GRPO 使用同一 prompt 的一组候选回答计算相对优势，不依赖单独的 value model。它适合具有可验证奖励的任务，例如格式是否正确、标签是否可抽取、标签是否与真实 label 一致。对于本项目，GRPO 可以作为 PPO 的轻量对照方案。")
    callout(doc, "实现建议", "LLaMA-Factory 主流程可覆盖 DPO、RM、PPO 等阶段；GRPO 是否直接可用取决于你安装的具体版本。若当前 LLaMA-Factory 不支持 GRPO，可使用 Hugging Face TRL、verl、MS-SWIFT 或支持 GRPO 的训练框架接入。", LIGHT_BLUE)
    table(doc, ["奖励项", "规则", "分值建议"], [
        ["标签正确", "预测标签与 metadata label 一致", "+1.0"],
        ["格式正确", "可抽取“故障类型”，full 模式字段齐全", "+0.3"],
        ["依据一致", "依据中描述的部位与故障类型匹配", "+0.2"],
        ["维护建议一致", "建议检查对象与故障类型匹配", "+0.2"],
        ["工况忠实", "不编造转速、负载、采样率", "+0.2"],
        ["幻觉惩罚", "出现不存在的故障类型或无关设备", "-0.5"],
    ], widths=[3.0, 8.4, 3.2], size=8.6)
    code(doc, """for prompt in batch:
    responses = policy.generate(prompt, num_return_sequences=4)
    rewards = [rule_reward(prompt, response, label) for response in responses]
    advantages = normalize_within_group(rewards)
    update_policy_with_grpo_loss(prompt, responses, advantages, kl_reference)

# 每个 prompt 至少生成 4 个候选，才能形成稳定组内相对比较。
# 奖励函数先用硬规则，后续可以融合 RM 分数。""", "GRPO 伪代码")
    code(doc, """{
  \"conversations\": [
    {\"from\": \"human\", \"value\": \"<image>\\n请诊断当前轴承状态...\"}
  ],
  \"label\": \"BF\",
  \"label_name\": \"滚动体故障\",
  \"condition\": \"20_0\",
  \"images\": [\"data/processed/images/all/sample_BF_20_0_000409.png\"]
}""", "GRPO prompt-only 数据示例")

    heading(doc, "8. 评估指标与实验矩阵", 1)
    table(doc, ["指标", "含义", "用途"], [
        ["Accuracy / Macro-F1", "从回答中抽取故障类型并与 label 比较", "保证 RL 不损伤分类能力"],
        ["格式正确率", "是否能稳定输出指定字段", "衡量对齐效果"],
        ["一致性通过率", "故障类型、依据、维护建议是否匹配", "衡量报告可信度"],
        ["幻觉率", "是否生成不存在的故障类型、工况或设备信息", "衡量安全性"],
        ["Preference Win Rate", "RL 模型回答优于 SFT 回答的比例", "评估偏好优化收益"],
        ["RM Pairwise Accuracy", "RM 是否给 chosen 更高分", "奖励模型质量门槛"],
        ["人工抽检通过率", "人工审核 50~100 条报告是否可接受", "最终 Demo 可信度"],
    ], widths=[3.2, 6.2, 5.0], size=8.5)
    table(doc, ["编号", "实验", "初始模型", "数据/奖励", "目标"], [
        ["R0", "SFT baseline", "checkpoint-2200", "class_only test", "固定基线"],
        ["R1", "DPO-hard", "checkpoint-2200", "规则构造 chosen/rejected", "验证偏好优化是否保分类"],
        ["R2", "DPO-mixed", "checkpoint-2200", "规则 + 模型生成 + 人工抽检偏好", "提升报告质量"],
        ["R3", "RM", "Qwen2.5-VL-3B", "偏好对", "训练 reward scorer"],
        ["R4", "PPO", "DPO adapter", "RM + 规则奖励", "在线优化完整报告"],
        ["R5", "GRPO", "DPO adapter", "多候选 + 规则/RM reward", "低 value-model 成本 RL 对照"],
    ], widths=[1.5, 3.0, 3.3, 4.5, 4.0], size=8.4)

    heading(doc, "9. 工程脚本与交付物规划", 1)
    table(doc, ["脚本", "功能", "输出"], [
        ["08_build_preference_data.py", "从 metadata、模型预测和规则模板生成 chosen/rejected 数据", "train_dpo.json、val_dpo.json"],
        ["09_validate_preference_data.py", "检查字段完整、标签一致、图片路径存在", "preference_data_report.json"],
        ["10_eval_reward_model.py", "评估 RM 对 chosen/rejected 的排序准确率", "rm_metrics.json"],
        ["11_eval_alignment.py", "计算格式正确率、一致性、幻觉率、偏好胜率", "alignment_metrics.json"],
        ["12_grpo_reward.py", "提供 GRPO/PPO 可复用的规则奖励函数", "reward function module"],
    ], widths=[4.2, 7.2, 4.0], size=8.4)
    code(doc, """configs/rlhf/
├── qwen2_5vl_3b_dpo_sufd.yaml
├── qwen2_5vl_3b_rm_sufd.yaml
├── qwen2_5vl_3b_ppo_sufd.yaml
└── qwen2_5vl_3b_grpo_sufd.yaml

data/processed/rlhf/
├── train_dpo.json
├── val_dpo.json
├── train_rm.json
├── rl_prompts.json
└── preference_data_report.json

outputs/rlhf/
├── dpo_metrics/
├── rm_metrics/
├── ppo_metrics/
└── grpo_metrics/""", "建议目录结构")

    heading(doc, "10. 风险控制与推荐执行顺序", 1)
    table(doc, ["风险", "表现", "控制策略"], [
        ["分类能力下降", "RL 后 Accuracy/Macro-F1 低于 checkpoint-2200", "降低学习率、缩短 epoch、增大 KL、回退 DPO"],
        ["奖励模型误导", "RM 给错误标签高分", "RM 未达验证集 pairwise accuracy 门槛不得用于 PPO"],
        ["模式坍塌", "模型总输出同一类或固定模板", "检查奖励分布，引入 KL 和多样化 prompts"],
        ["偏好数据噪声", "chosen/rejected 本身不可靠", "强规则构造 + 人工抽检 100 条"],
        ["显存不足", "PPO/GRPO 训练 OOM", "降低 batch、使用 QLoRA、先 DPO，PPO 小样本"],
        ["GRPO 框架不兼容", "当前 LLaMA-Factory 无 GRPO stage", "使用 TRL/verl/MS-SWIFT 或只作为后续扩展"],
    ], widths=[3.0, 4.6, 7.0], size=8.4)
    table(doc, ["优先级", "任务", "完成标准"], [
        ["P0", "构造 1000~3000 条高质量偏好对", "图片路径存在，chosen/rejected 标签和格式可校验"],
        ["P1", "先跑 DPO-hard", "分类指标不低于 0.48 太多，格式正确率保持 1.0"],
        ["P2", "训练 RM", "验证集 pairwise accuracy 建议 ≥ 0.75"],
        ["P3", "PPO 小样本试跑", "reward 上升但 KL、长度、分类指标稳定"],
        ["P4", "GRPO 对照", "若框架可用，比较 GRPO 与 PPO/DPO 的成本和收益"],
        ["P5", "Demo 升级", "Demo 展示 RLHF 版报告质量对比"],
    ], widths=[2.0, 6.4, 7.0], size=8.5)
    callout(doc, "最终建议", "先实现 DPO，不要直接从 PPO 或 GRPO 开始。DPO 的数据、训练和评估成本最低，能最快判断“偏好对齐”是否能提升诊断报告质量且不显著牺牲分类能力。RM/PPO/GRPO 应作为增强版路线逐步推进。", LIGHT_GREEN)

    heading(doc, "参考资料与实现依据", 1)
    bullets(doc, [
        "LLaMA-Factory Documentation: https://llamafactory.readthedocs.io/",
        "LLaMA-Factory Data Preparation: https://llamafactory.readthedocs.io/en/latest/getting_started/data_preparation.html",
        "LLaMA-Factory Advanced Arguments: https://www.aidoczh.com/llamafactory/en/advanced/arguments.html",
        "Hugging Face TRL GRPO Trainer: https://huggingface.co/docs/trl/en/grpo_trainer",
        "Qwen2.5-VL-3B-Instruct model card: https://huggingface.co/Qwen/Qwen2.5-VL-3B-Instruct",
    ])

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run("SUFD-QwenVL RLHF Project Plan")
        qfont(r, size=8, color="808080")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
